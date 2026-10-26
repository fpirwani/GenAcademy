"""Generate a small, synthetic version pair; does not upload or index user documents."""

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def create_samples():
    folder = ROOT / "samples"
    folder.mkdir(exist_ok=True)
    for version, notice in [(1, "30 days"), (2, "15 business days")]:
        doc = Document()
        doc.add_heading("Example Service Policy", 0)
        doc.add_heading("Cancellation", 1)
        doc.add_paragraph(
            f"The customer must provide at least {notice} written notice before cancellation."
        )
        doc.add_heading("Payment", 1)
        doc.add_paragraph("Invoices must be paid within 30 days of receipt.")
        if version == 2:
            doc.add_heading("Renewal", 1)
            doc.add_paragraph(
                "The agreement renews automatically each year unless cancelled in writing."
            )
        doc.save(folder / f"policy-v{version}.docx")
    return folder


if __name__ == "__main__":
    print(create_samples())
