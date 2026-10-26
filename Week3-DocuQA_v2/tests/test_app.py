from pathlib import Path

from streamlit.testing.v1 import AppTest
from test_documents import pdf_bytes

import config
import models
from documents import extract, make_chunks
from storage import Store


def setup_app(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(models, "model_status", lambda: (False, "Start Ollama"))
    store = Store(tmp_path)
    metadata, _ = store.add("planets.pdf", pdf_bytes())
    blocks = extract("planets.pdf", store.original(metadata["id"]).read_bytes())
    store.write(metadata["id"], metadata["generation"], "blocks", blocks)
    store.write(metadata["id"], metadata["generation"], "chunks", make_chunks(blocks))
    store.write(
        metadata["id"], metadata["generation"], "summary", {"text": "A report about planets."}
    )
    store.update(metadata["id"], metadata["generation"], status="ready")
    return store, metadata


def test_latest_message_last_without_sample_questions(tmp_path, monkeypatch):
    store, doc = setup_app(tmp_path, monkeypatch)
    history = [
        {"role": "user", "content": "Old question"},
        {"role": "assistant", "content": "Old answer"},
        {"role": "user", "content": "New question"},
        {"role": "assistant", "content": "New answer"},
    ]
    store.write(doc["id"], doc["generation"], "chat", history)
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py")).run()
    assert not app.exception
    assert not [button for button in app.button if str(button.key).startswith("sample_")]
    assert [message.text[0].value for message in app.chat_message] == [
        "Old question",
        "Old answer",
        "New question",
        "New answer",
    ]
    assert store.read(doc["id"], "chat") == history


def test_empty_app_loads_without_ollama(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(models, "model_status", lambda: (False, "Start Ollama"))
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py")).run()
    assert not app.exception
    assert any("DocuQA" in title.value for title in app.title)


def test_saved_summary_can_be_regenerated_as_a_plain_paragraph(tmp_path, monkeypatch):
    from test_workflows import FakeModel

    import tools

    store, doc = setup_app(tmp_path, monkeypatch)
    monkeypatch.setattr(models, "model_status", lambda: (True, "Local model ready"))
    original = tools.build_tools
    monkeypatch.setattr(
        tools,
        "build_tools",
        lambda *args, **kwargs: original(*args, **kwargs, model_factory=FakeModel),
    )
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py")).run()
    assert not app.exception
    assert any(panel.label == "Summary" for panel in app.expander)
    next(button for button in app.button if button.label == "Regenerate summary").click().run()
    assert not app.exception
    assert "Mercury" in store.read(doc["id"], "summary")["text"]


def test_interrupted_indexing_can_be_retried(tmp_path, monkeypatch):
    store, doc = setup_app(tmp_path, monkeypatch)
    (store.folder(doc["id"]) / "chunks.json").unlink()
    store.update(doc["id"], doc["generation"], status="uploaded")
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py")).run()
    assert not app.exception
    retry = next((button for button in app.button if button.label == "Retry file processing"), None)
    assert retry is not None, "Incomplete search index must offer recovery"
    retry.click().run()
    assert not app.exception
    assert len(store.read(doc["id"], "chunks")) == 2


def test_source_click_navigates_preview_and_delete_removes_document(tmp_path, monkeypatch):
    store, doc = setup_app(tmp_path, monkeypatch)
    store.write(
        doc["id"],
        doc["generation"],
        "chat",
        [
            {"role": "user", "content": "What about Venus?"},
            {
                "role": "assistant",
                "content": "Venus has a thick atmosphere.",
                "sources": [
                    {
                        "chunk_id": "c2",
                        "location": 2,
                        "label": "Page 2",
                        "quote": "Venus has a thick atmosphere.",
                    }
                ],
            },
        ],
    )
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py")).run()
    assert not app.exception
    source = next(button for button in app.button if button.label == "Page 2")
    source.click().run()
    assert not app.exception
    assert app.number_input[0].value == 2
    next(button for button in app.button if button.label == "Delete file").click().run()
    next(button for button in app.button if button.label == "Delete permanently").click().run()
    assert not app.exception
    assert store.list() == []


def test_docx_source_navigation_and_document_switch_isolate_chat(tmp_path, monkeypatch):
    from io import BytesIO

    from docx import Document

    store, pdf = setup_app(tmp_path, monkeypatch)
    word = Document()
    word.add_paragraph("Introduction to the budget.")
    word.add_page_break()
    word.add_paragraph("The equipment budget is $400.")
    payload = BytesIO()
    word.save(payload)
    doc, _ = store.add("budget.docx", payload.getvalue())
    blocks = extract("budget.docx", payload.getvalue())
    for name, value in [
        ("blocks", blocks),
        ("chunks", make_chunks(blocks)),
        ("summary", {"text": "Budget overview."}),
    ]:
        store.write(doc["id"], doc["generation"], name, value)
    store.write(
        doc["id"],
        doc["generation"],
        "chat",
        [
            {
                "role": "assistant",
                "content": "The equipment budget is $400.",
                "sources": [
                    {
                        "chunk_id": "c2",
                        "location": 2,
                        "label": "Block 2",
                        "quote": "The equipment budget is $400.",
                    }
                ],
            }
        ],
    )
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py")).run()
    next(button for button in app.button if button.label == "Block 2").click().run()
    assert not app.exception
    assert app.number_input[0].label == "Page"
    assert app.number_input[0].value == 2
    assert len(app.chat_message) == 1
    app.radio(key="select_file_" + pdf["id"]).set_value(pdf["id"]).run()
    assert not app.exception
    assert len(app.chat_message) == 0
    assert app.number_input[0].value == 1


def test_delete_confirmation_does_not_transfer_to_reuploaded_file(tmp_path, monkeypatch):
    store, doc = setup_app(tmp_path, monkeypatch)
    payload = store.original(doc["id"]).read_bytes()
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py")).run()
    next(button for button in app.button if button.label == "Delete file").click().run()
    confirm = next(button for button in app.button if button.label == "Delete permanently")
    store.delete(doc["id"])
    replacement, _ = store.add("planets.pdf", payload)
    confirm.click().run()
    assert not app.exception
    assert store.read(doc["id"], "metadata")["generation"] == replacement["generation"]
