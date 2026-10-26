from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas

from documents import extract, make_chunks, render_page


def pdf_bytes():
    output = BytesIO()
    canvas = Canvas(output)
    canvas.drawString(30, 700, "Mercury is the closest planet to the Sun.")
    canvas.showPage()
    canvas.drawString(30, 700, "Venus has a thick atmosphere.")
    canvas.save()
    return output.getvalue()


def test_pdf_preserves_page_sources_and_renders():
    payload = pdf_bytes()
    blocks = extract("planets.pdf", payload)
    assert [b["location"] for b in blocks] == [1, 2]
    assert "Mercury" in blocks[0]["text"]
    chunks = make_chunks(blocks)
    assert chunks[1]["location"] == 2
    assert render_page(payload, 2).startswith(b"\x89PNG")


def test_docx_includes_tables_in_reading_order():
    doc = Document()
    doc.add_heading("Costs", level=1)
    doc.add_paragraph("The budget follows.")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Equipment: $400"
    doc.add_paragraph("End of report.")
    output = BytesIO()
    doc.save(output)
    blocks = extract("budget.docx", output.getvalue())
    assert "Costs" in blocks[0]["text"]
    assert "Equipment: $400" in blocks[2]["text"]
    assert blocks[3]["location"] == 4


def test_empty_encrypted_invalid_and_oversize_files_are_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="text"):
        extract("blank.pdf", output.getvalue())
    writer.encrypt("secret")
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="password|encrypted"):
        extract("locked.pdf", output.getvalue())
    with pytest.raises(ValueError):
        extract("wrong.doc", b"not a docx")
    with pytest.raises(ValueError):
        extract("broken.pdf", b"bad")
    with pytest.raises(ValueError, match="25"):
        extract("large.pdf", b"a" * (25 * 1024 * 1024 + 1))
